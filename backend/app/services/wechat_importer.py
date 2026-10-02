"""微信支付账单流水 Importer（xlsx 格式）。

真实账单结构（Sheet1，11 列）：
交易时间 | 交易类型 | 交易对方 | 商品 | 收/支 | 金额(元) | 支付方式 | 当前状态 | 交易单号 | 商户单号 | 备注

设计约定（与银行/支付宝 Importer 同一契约）：
- matches_file()：xlsx 魔数（PK\\x03\\x04）+ openpyxl 读取首表头 → detect
- parse()：单元格统一转字符串存 Raw（datetime → 'YYYY-MM-DD HH:MM:SS'，金额数值 → str，
  '/​' 与全横线占位符归一为空串），Raw 不可变
- normalize()：amount 为正数幅值，方向由 收/支 列表达（微信金额不带符号）
- 微信账单自带全局唯一「交易单号」→ source_transaction_id 直接使用，无需合成
- 支付方式为 '/' 时（如零钱收入行），用「当前状态」文本兜底（「已存入零钱」→
  含「零钱」关键词，可匹配零钱账户别名）
"""
from typing import Dict, List

import os
import shutil
import tempfile

from app.core.utils import parse_decimal
from app.services.importer import BaseImporter, ImporterRegistry, NormalizedTransaction

_OPENPYXL_SUFFIXES = (".xlsx", ".xlsm", ".xltx", ".xltm")


def _is_xlsx(file_path: str) -> bool:
    try:
        with open(file_path, "rb") as f:
            return f.read(4) == b"PK\x03\x04"
    except OSError:
        return False


def _ensure_xlsx_readable(file_path: str) -> str:
    """openpyxl 按文件扩展名校验格式；storage 层保存的 temp 文件无扩展名
    （temp_<uuid>），会直接抛 InvalidFileException。这里把无扩展名文件复制到
    带 .xlsx 后缀的临时路径再打开（账单文件 <1MB，复制成本可忽略）。
    返回可打开的路径；调用方负责在 path != 原路径时删除临时副本。"""
    if file_path.lower().endswith(_OPENPYXL_SUFFIXES):
        return file_path
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    os.close(fd)
    shutil.copyfile(file_path, tmp)
    return tmp


def _cleanup_xlsx_readable(path: str, original: str):
    if path != original and os.path.exists(path):
        os.remove(path)


def _clean_cell(value) -> str:
    """xlsx 单元格 → 字符串：datetime → 'YYYY-MM-DD HH:MM:SS'，数值 → str，
    '/' 与全横线占位符归一为空串。金额经 str() 转换保持与账面一致（Python str
    取 float 最短表示，如 9.82），下游 parse_decimal 全程 Decimal。"""
    if value is None:
        return ""
    s = str(value).strip()
    if s in ("", "/") or (s and set(s) <= {"-", "/"}):
        return ""
    return s


class WeChatImporter(BaseImporter):
    name = "wechat"
    COLUMNS = ["交易时间", "交易类型", "交易对方", "商品", "收/支", "金额(元)",
               "支付方式", "当前状态", "交易单号", "商户单号", "备注"]
    REQUIRED_COLUMNS = {"交易时间", "交易对方", "收/支", "金额(元)", "支付方式", "交易单号"}

    def detect(self, raw_data: Dict) -> bool:
        keys = set(raw_data.keys())
        return self.REQUIRED_COLUMNS <= keys

    def matches_file(self, file_path: str) -> bool:
        if not _is_xlsx(file_path):
            return False
        import openpyxl
        path = _ensure_xlsx_readable(file_path)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                # 扫描前 10 行找表头（兼容带元数据前导行的导出格式）
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 10:
                        break
                    probe = {_clean_cell(c): "" for c in (row or [])}
                    if self.detect(probe):
                        return True
                return False
            finally:
                wb.close()
        except Exception:
            return False
        finally:
            _cleanup_xlsx_readable(path, file_path)

    def parse(self, file_path: str) -> List[Dict]:
        import openpyxl
        rows: List[Dict] = []
        path = _ensure_xlsx_readable(file_path)
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            try:
                ws = wb[wb.sheetnames[0]]
                for idx, row in enumerate(ws.iter_rows(values_only=True)):
                    cells = [_clean_cell(c) for c in (row or [])]
                    if not cells or not cells[0]:
                        continue
                    if cells[0] == "交易时间":
                        continue  # 表头行
                    d = dict(zip(self.COLUMNS, cells))
                    if not d.get("交易单号"):
                        continue  # 无交易单号的行不可信，跳过
                    rows.append(d)
            finally:
                wb.close()
        finally:
            _cleanup_xlsx_readable(path, file_path)
        return rows

    def normalize(self, raw_data: Dict) -> NormalizedTransaction:
        ts = raw_data.get("交易时间") or ""
        date_part, _, time_part = ts.partition(" ")
        direction = raw_data.get("收/支") or "不计收支"
        payment_method = raw_data.get("支付方式") or ""
        if not payment_method:
            # 零钱收入行等场景：支付方式列为 '/'，资金去向在「当前状态」（如「已存入零钱」）
            payment_method = raw_data.get("当前状态") or ""
        return NormalizedTransaction(
            source_transaction_id=raw_data.get("交易单号") or "",
            date=date_part,
            time=time_part,
            amount=parse_decimal(raw_data.get("金额(元)")),
            currency="CNY",
            direction=direction,
            transaction_type=raw_data.get("交易类型") or "",
            merchant=raw_data.get("交易对方") or "",
            description=raw_data.get("商品") or "",
            payment_method=payment_method,
            counterparty=raw_data.get("交易对方") or "",
        )


ImporterRegistry.register(WeChatImporter())
